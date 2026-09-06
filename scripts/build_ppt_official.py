# -*- coding: utf-8 -*-
"""
SatQuery AI - SIH 2026 Idea Submission.

Builds onto the OFFICIAL SIH 2026 template so every piece of official chrome
(logo, team oval, blue footer bar, title placeholders, slide numbers) is the
template's own, untouched. Only the content is ours.

Run:  .venv/Scripts/python.exe scripts/build_ppt_official.py
"""
import copy
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = r"C:\Users\Ariha\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
OUT = os.path.join(ROOT, "SatQueryAI-SIH2026-Idea-Submission.pptx")

TEAM_NAME = "Oryn"
TEAM_ID = "Not assigned yet"
PS_ID = "26167"
PS_TITLE = ("SatQuery AI: An Interactive Vision-Language Assistant for Multimodal "
            "Remote Sensing Image Analysis through Text Queries")

# --- colours sampled from the template itself ------------------------------
NAVY = RGBColor(0x1F, 0x4E, 0x7D)
INK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
ARIAL = "Arial"

# diagram fills
F_IN = "FFF2CC"
F_PROC = "DEEBF7"
F_DEC = "E2EFDA"
F_REG = "FBE5D6"
F_OUT = "D9E2F3"

BODY_TOP = 1.30
BODY_BOT = 6.88


# ---------------------------------------------------------------- utilities
def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def find(slide, predicate):
    for sh in slide.shapes:
        if predicate(sh):
            return sh
    return None


def by_name(slide, prefix):
    return find(slide, lambda s: s.name.startswith(prefix))


def drop(shape):
    shape._element.getparent().remove(shape._element)


def set_oval(slide):
    ov = find(slide, lambda s: s.name.startswith("Oval"))
    if ov is None:
        return
    tf = ov.text_frame
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        p.runs[0].text = TEAM_NAME
    else:
        p.add_run().text = TEAM_NAME
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.name = ARIAL


def set_title(slide, text, size=32):
    ttl = by_name(slide, "Title")
    if ttl is None:
        return
    tf = ttl.text_frame
    p = tf.paragraphs[0]
    runs = list(p.runs)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = text
        runs[0].font.size = Pt(size)
    else:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = True


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, first, parts, size=11, space_after=5, bullet=False,
         color=INK, space_before=0, align=None, line=None):
    """parts = [(text, bold), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if bullet:
        r = p.add_run()
        r.text = "\u2022  "
        r.font.size = Pt(size)
        r.font.name = ARIAL
        r.font.color.rgb = NAVY
        r.font.bold = True
    for txt, bold in parts:
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bool(bold)
        r.font.name = ARIAL
        r.font.color.rgb = color
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if align:
        p.alignment = align
    if line:
        p.line_spacing = line
    return p


def section(slide, text, x, y, w, size=13, underline=True):
    tb, tf = textbox(slide, x, y, w, 0.30)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = ARIAL
    r.font.color.rgb = NAVY
    r.font.underline = underline
    return tb


def vline(slide, x, y1, y2, color="BFBFBF"):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y1),
                                Inches(0.008), Inches(y2 - y1))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGBColor.from_string(color)
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln


def hline(slide, x1, x2, y, color="BFBFBF"):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x1), Inches(y),
                                Inches(x2 - x1), Inches(0.008))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGBColor.from_string(color)
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln


def node(slide, text, x, y, w, h, fill, size=8.5, bold=False, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.color.rgb = RGBColor.from_string("808080")
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for i, ln_ in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln_
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = ARIAL
        r.font.color.rgb = INK
    return sh


def arrow(slide, cx, y, h):
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - 0.006),
                                  Inches(y), Inches(0.012), Inches(h - 0.05))
    stem.fill.solid()
    stem.fill.fore_color.rgb = RGBColor.from_string("808080")
    stem.line.fill.background()
    stem.shadow.inherit = False
    tip = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx - 0.05),
                                 Inches(y + h - 0.07), Inches(0.10), Inches(0.08))
    tip.rotation = 180
    tip.fill.solid()
    tip.fill.fore_color.rgb = RGBColor.from_string("808080")
    tip.line.fill.background()
    tip.shadow.inherit = False


# ---------------------------------------------------------------- build
prs = Presentation(TEMPLATE)

# drop the "IMPORTANT INSTRUCTIONS" slide (template says max 6 slides)
delete_slide(prs, 6)

S = prs.slides

# ============================ SLIDE 1 - TITLE PAGE ========================
s1 = S[0]
tb = by_name(s1, "TextBox 9")
tf = tb.text_frame
fields = [
    ("Problem Statement ID - ", PS_ID, 17),
    ("Problem Statement Title - ", PS_TITLE, 13),
    ("Theme - ", "Space Technology", 17),
    ("PS Category - ", "Software", 17),
    ("Team ID - ", TEAM_ID, 17),
    ("Team Name - ", TEAM_NAME, 17),
]
paras = list(tf.paragraphs)
for i, (label, value, size) in enumerate(fields):
    p = paras[i + 1]                      # paras[0] is the empty spacer
    runs = list(p.runs)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = label
        runs[0].font.size = Pt(size)
        runs[0].font.bold = True
        runs[0].font.name = ARIAL
    r = p.add_run()
    r.text = value
    r.font.size = Pt(size)
    r.font.bold = False
    r.font.name = ARIAL
    p.space_after = Pt(10)

# ====================== SLIDE 2 - PROPOSED SOLUTION =======================
s2 = S[1]
set_oval(s2)
set_title(s2, "SatQuery AI", 34)
drop(by_name(s2, "TextBox 8"))

MID_Y = 2.62
vline(s2, 6.30, BODY_TOP, MID_Y - 0.10)
hline(s2, 0.40, 12.95, MID_Y - 0.04)

section(s2, "Problem", 0.40, BODY_TOP, 3.0)
tb, tf = textbox(s2, 0.40, BODY_TOP + 0.30, 5.75, 1.05)
para(tf, True, [("Satellite imagery is still read by hand.", True),
                (" An analyst with a Cartosat-2S optical or RISAT SAR scene must open a "
                 "GIS desktop and interpret it manually. Generic AI models cannot help: "
                 "they ", False),
                ("hallucinate on nadir imagery, cannot read multi-band GeoTIFF, and are "
                 "blind to SAR backscatter", True), (".", False)],
     size=11.5, line=1.15)

section(s2, "Our Idea", 6.55, BODY_TOP, 3.0)
tb, tf = textbox(s2, 6.55, BODY_TOP + 0.30, 6.40, 1.05)
para(tf, True, [("SatQuery AI", True),
                (" answers plain-English questions about satellite imagery - a single "
                 "scene, a ", False), ("before-and-after pair", True), (", or a "
                 "co-registered ", False), ("optical + SAR pair", True),
                (". An agentic controller routes each query to remote-sensing-adapted "
                 "models and returns the answer with an ", False),
                ("auditable execution trace.", True)], size=11.5, line=1.15)

# --- row 2
section(s2, "Proposed Solution", 0.40, MID_Y + 0.06, 4.2)
tb, tf = textbox(s2, 0.40, MID_Y + 0.40, 4.15, 4.0)
rows = [
    [("An agentic controller ", True),
     ("interprets the query, classifies the task, validates the images, selects models "
      "from a typed registry, executes them, fuses text with spatial output and scores "
      "confidence.", False)],
    [("Remote-sensing adaptation: ", True),
     ("a 12-channel encoder fine-tuned on BigEarthNet consumes 10 optical and 2 SAR bands "
      "together, so fusion is learned, not stapled on afterwards.", False)],
    [("Compositional change analysis: ", True),
     ("a siamese detector produces a change mask, statistics are derived from it, and the "
      "language model verbalises them - every step is auditable.", False)],
    [("Geospatial correctness first: ", True),
     ("CRS, ground sample distance and sub-pixel grid alignment are verified before any "
      "cross-modal claim.", False)],
]
for i, parts in enumerate(rows):
    para(tf, i == 0, parts, size=10.5, bullet=True, space_after=7, line=1.12)

# --- centre diagram
dx, dw = 4.78, 3.72
tb, tf = textbox(s2, dx, MID_Y + 0.40, dw, 0.22)
para(tf, True, [("HOW A QUERY IS ANSWERED", True)], size=9, color=NAVY,
     align=PP_ALIGN.CENTER, space_after=0)

y = MID_Y + 0.70
steps = [
    ('"What changed here?"  +  two GeoTIFFs', F_IN, 0.32),
    ("INGEST - CRS, bands, GSD, modality", F_PROC, 0.32),
    ("VALIDATE - format, overlap, grid alignment", F_PROC, 0.32),
    ("CLASSIFY task  \u2192  SELECT model by rule", F_DEC, 0.32),
]
for txt, fill, hh in steps:
    node(s2, txt, dx, y, dw, hh, fill, size=8.5)
    y += hh
    arrow(s2, dx + dw / 2, y, 0.12)
    y += 0.12

reg = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(dx), Inches(y),
                          Inches(dw), Inches(0.76))
reg.fill.solid()
reg.fill.fore_color.rgb = RGBColor.from_string(F_REG)
reg.line.color.rgb = RGBColor.from_string("808080")
reg.line.width = Pt(0.75)
reg.shadow.inherit = False
rtf = reg.text_frame
rtf.margin_top = Inches(0.03)
rtf.vertical_anchor = MSO_ANCHOR.TOP
rp = rtf.paragraphs[0]
rp.alignment = PP_ALIGN.CENTER
rr = rp.add_run()
rr.text = "MODEL REGISTRY"
rr.font.size = Pt(8)
rr.font.bold = True
rr.font.name = ARIAL
rr.font.color.rgb = INK
rw = (dw - 0.24) / 3
for i, t in enumerate(["Optical+SAR\nEncoder", "Change\nDetector",
                       "RS Vision-\nLanguage Model"]):
    node(s2, t, dx + 0.06 + i * (rw + 0.06), y + 0.22, rw, 0.48, "FFFFFF", size=7)
y += 0.76
arrow(s2, dx + dw / 2, y, 0.12)
y += 0.12
node(s2, "FUSE text + spatial  →  CONFIDENCE", dx, y, dw, 0.30, F_DEC, size=8.5)
y += 0.30
arrow(s2, dx + dw / 2, y, 0.12)
y += 0.12
node(s2, "ANSWER  +  AUDITABLE TRACE", dx, y, dw, 0.32, F_OUT, size=8.5, bold=True)
y += 0.32
assert y <= BODY_BOT, "slide2 diagram overflow: ends at %.2f, budget %.2f" % (y, BODY_BOT)

section(s2, "Innovation & Uniqueness", 8.75, MID_Y + 0.06, 4.3)
tb, tf = textbox(s2, 8.75, MID_Y + 0.40, 4.20, 4.0)
rows = [
    [("Deterministic model selection. ", True),
     ("The language model classifies the task; typed registry rules choose the model, so "
      "the trace is reproducible - not a chat transcript.", False)],
    [("One fused optical+SAR encoder", True),
     (", trained on co-registered Sentinel-1/2 pairs and proven by a three-way ablation "
      "where fused beats either sensor alone.", False)],
    [("Co-registration is verified, not assumed. ", True),
     ("Sub-pixel grid alignment is checked, not merely footprint overlap.", False)],
    [("Confidence as a named breakdown", True),
     (", including agreement between the language model and the specialist - disagreement "
      "visibly lowers the score.", False)],
    [("Sensor-correct preprocessing. ", True),
     ("SAR is converted to decibels and percentile-clipped, so the model sees terrain "
      "rather than a black frame.", False)],
]
for i, parts in enumerate(rows):
    para(tf, i == 0, parts, size=10.5, bullet=True, space_after=6, line=1.12)

# ====================== SLIDE 3 - TECHNICAL APPROACH ======================
s3 = S[2]
set_oval(s3)
set_title(s3, "TECHNICAL APPROACH", 32)
drop(by_name(s3, "TextBox 8"))

SPLIT = 4.45
vline(s3, SPLIT, BODY_TOP, BODY_BOT)

section(s3, "Technologies Used", 0.40, BODY_TOP, 3.8)
tb, tf = textbox(s3, 0.40, BODY_TOP + 0.30, 3.85, 5.2)
tech = [
    [("Python 3.11 + FastAPI: ", True),
     ("async backend with native server-sent events, so the execution trace streams as "
      "each step completes.", False)],
    [("LangGraph: ", True),
     ("typed state machine with conditional edges. Chosen over a plain agent loop because "
      "the trace needs inspectable, reproducible transitions.", False)],
    [("rasterio + GDAL: ", True),
     ("reads multi-band GeoTIFF with CRS, affine transform and nodata. Chosen over "
      "OpenCV/PIL, which discard geospatial metadata.", False)],
    [("pyproj + Shapely: ", True),
     ("projection comparison and footprint intersection for co-registration checks.", False)],
    [("PyTorch: ", True),
     ("12-channel ResNet-50 encoder fine-tuned on BigEarthNet v2.0 for 19-class land "
      "cover.", False)],
    [("Qwen2.5-VL (RS-adapted) + LoRA: ", True),
     ("3B vision-language model. Chosen over 7B because domain knowledge sits in the "
      "encoder - same answer, commodity hardware.", False)],
    [("React + Vite: ", True), ("live trace panel driven by the SSE stream.", False)],
    [("Kaggle T4: ", True), ("model adaptation runs off the demo machine entirely.", False)],
]
for i, parts in enumerate(tech):
    para(tf, i == 0, parts, size=9.5, bullet=True, space_after=6, line=1.10)

section(s3, "Methodology - Flow Chart", SPLIT + 0.25, BODY_TOP, 4.5)

fx = SPLIT + 0.25
fw = 12.95 - fx
y = BODY_TOP + 0.34
iw = (fw - 0.20) / 3
for i, t in enumerate(["Text query", "Optical GeoTIFF\n(Cartosat-2S)",
                       "SAR / second-date GeoTIFF\n(RISAT)"]):
    node(s3, t, fx + i * (iw + 0.10), y, iw, 0.40, F_IN, size=8.5)
y += 0.40
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18

node(s3, "INGEST  -  checksum, CRS, transform, bands, dtype, nodata, modality from pixels",
     fx, y, fw, 0.34, F_PROC, size=9)
y += 0.34
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18

node(s3, "VALIDATE  -  format, projection match, footprint overlap, GSD ratio, sub-pixel "
         "grid alignment", fx, y, fw, 0.34, F_PROC, size=9)
y += 0.34
tb, tf = textbox(s3, fx + fw - 2.75, y + 0.01, 2.70, 0.18)
para(tf, True, [("fatal  \u2192  reject with a named reason", False)], size=7.5,
     color=RED, align=PP_ALIGN.RIGHT, space_after=0)
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18

cw = (fw - 0.10) / 2
node(s3, "CLASSIFY TASK\n(language model, structured output)", fx, y, cw, 0.42, F_DEC,
     size=8.5)
node(s3, "PLAN - REGISTRY RESOLVER\n(deterministic rules, logged rejections)",
     fx + cw + 0.10, y, cw, 0.42, F_DEC, size=8.5)
y += 0.42
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18

reg = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx), Inches(y), Inches(fw),
                          Inches(0.82))
reg.fill.solid()
reg.fill.fore_color.rgb = RGBColor.from_string(F_REG)
reg.line.color.rgb = RGBColor.from_string("808080")
reg.line.width = Pt(0.75)
reg.shadow.inherit = False
rtf = reg.text_frame
rtf.margin_top = Inches(0.03)
rtf.vertical_anchor = MSO_ANCHOR.TOP
rp = rtf.paragraphs[0]
rp.alignment = PP_ALIGN.CENTER
rr = rp.add_run()
rr.text = "EXECUTE  -  model registry"
rr.font.size = Pt(8.5)
rr.font.bold = True
rr.font.name = ARIAL
rr.font.color.rgb = INK
mw = (fw - 0.28) / 3
for i, t in enumerate([
        "12-channel Optical+SAR Encoder\nfine-tuned on BigEarthNet v2.0",
        "Siamese Change Detector\nmask + change statistics",
        "RS Vision-Language Model\nVQA, captioning, grounding"]):
    node(s3, t, fx + 0.07 + i * (mw + 0.07), y + 0.22, mw, 0.54, "FFFFFF", size=7)
y += 0.82
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18

node(s3, "FUSE text + geometry   \u2192   CONFIDENCE (input quality, task certainty, model "
         "fitness, cross-model agreement)", fx, y, fw, 0.34, F_DEC, size=8.5)
y += 0.34
arrow(s3, fx + fw / 2, y, 0.18)
y += 0.18
node(s3, "ANSWER  +  BOX/MASK OVERLAY  +  AUDITABLE EXECUTION TRACE", fx, y, fw, 0.36,
     F_OUT, size=9, bold=True)
y += 0.36

strip = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx), Inches(y + 0.12),
                            Inches(fw), Inches(BODY_BOT - (y + 0.14)))
strip.fill.solid()
strip.fill.fore_color.rgb = RGBColor.from_string(F_OUT)
strip.line.color.rgb = RGBColor.from_string("9DC3E6")
strip.line.width = Pt(0.75)
strip.shadow.inherit = False
stf = strip.text_frame
stf.word_wrap = True
stf.margin_left = stf.margin_right = Inches(0.1)
stf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = stf.paragraphs[0]
for txt, bold in [("Training & evaluation data: ", True),
                  ("BigEarthNet v2.0 (adaptation)  \u00b7  VRSBench + RSVQA (captioning, "
                   "grounding, VQA)  \u00b7  CDVQA (change VQA)", False)]:
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(9)
    r.font.bold = bold
    r.font.name = ARIAL
    r.font.color.rgb = INK
p2 = stf.add_paragraph()
for txt, bold in [("GitHub: ", True), ("[repository link]        ", False),
                  ("Demo video: ", True), ("[video link]", False)]:
    r = p2.add_run()
    r.text = txt
    r.font.size = Pt(9)
    r.font.bold = bold
    r.font.name = ARIAL
    r.font.color.rgb = INK

# ================= SLIDE 4 - FEASIBILITY AND VIABILITY ====================
s4 = S[3]
set_oval(s4)
set_title(s4, "FEASIBILITY AND VIABILITY", 32)
drop(by_name(s4, "TextBox 8"))

MID_X, ROW_Y = 6.68, 4.16
vline(s4, MID_X, BODY_TOP, BODY_BOT)
hline(s4, 0.40, MID_X - 0.05, ROW_Y - 0.10)
hline(s4, MID_X + 0.05, 12.95, ROW_Y - 0.10)

quad = [
    ("Feasibility", 0.40, BODY_TOP, [
        [("Every dataset the problem statement names is publicly downloadable", True),
         (" and licence-verified - no gated access, no institutional request.", False)],
        [("BigEarthNet v2.0 provides 549,488 co-registered Sentinel-1/2 pairs", True),
         (" - the only large public archive of aligned optical and radar imagery, so "
          "cross-modal adaptation needs no proprietary data.", False)],
        [("Adaptation fits free compute. ", True),
         ("A 12-channel encoder plus LoRA trains on a free Kaggle T4 - no cluster, no "
          "procurement.", False)],
        [("The geospatial stack installs cleanly. ", True),
         ("rasterio ships GDAL-bundled wheels, removing the usual build barrier.", False)],
        [("Hardware-independent by design. ", True),
         ("The model layer is swappable, so the same system runs on a 4 GB laptop GPU or "
          "Apple Silicon.", False)],
    ]),
    ("Commercial & Operational Viability", MID_X + 0.25, BODY_TOP, [
        [("Market validation: ", True),
         ("ISRO/NRSC's Bhoonidhi portal already disseminates open data from 47 satellites. "
          "The bottleneck is analyst time, not imagery.", False)],
        [("Technical readiness: high. ", True),
         ("Every component has a published open-source precedent; the work is integration "
          "and adaptation, not invention.", False)],
        [("Operational cost: low. ", True),
         ("Commodity hardware, open weights, open datasets - no per-seat licence, no GPU "
          "cluster.", False)],
        [("Sustainable to extend: ", True),
         ("new sensors and models are added as registry configuration, not new code.", False)],
        [("Deployment path: ", True),
         ("a self-contained web service that can sit inside an institutional network, since "
          "no imagery has to leave the premises.", False)],
    ]),
    ("Potential Challenges & Risks", 0.40, ROW_Y, [
        [("SAR has extreme dynamic range. ", True),
         ("A naive contrast stretch yields a near-black image, and a model will then "
          "confidently describe an empty scene.", False)],
        [("Real ISRO scenes are not benchmark tiles. ", True),
         ("Band order, data type, nodata convention and scene size all differ from public "
          "512-pixel datasets.", False)],
        [("CDVQA ships question text only", True),
         (" - the change imagery must be sourced separately from the SECOND dataset.", False)],
        [("Generic vision-language models hallucinate", True),
         (" on nadir imagery and cannot be trusted unsupervised.", False)],
        [("Co-registration is easy to claim and hard to prove", True),
         (" from footprint overlap alone.", False)],
    ]),
    ("Strategies for Overcoming Them", MID_X + 0.25, ROW_Y, [
        [("Sensor-correct preprocessing: ", True),
         ("decibel conversion with 2nd/98th percentile clipping and dual-polarisation false "
          "colour, validated against real EOS-04 scenes.", False)],
        [("Scene-scale ingestion: ", True),
         ("windowed reads and tiling, with the tiling recorded in the trace so aggregated "
          "answers stay explainable.", False)],
        [("Mirror the change dataset early", True),
         (" and cross-check its filenames against the question set before depending on it.",
          False)],
        [("Ground every claim in a specialist model. ", True),
         ("Cross-model disagreement automatically lowers reported confidence.", False)],
        [("Verify sub-pixel grid alignment", True),
         (", not just overlap, before any cross-modal answer is produced.", False)],
    ]),
]
for title, x, y0, rows in quad:
    section(s4, title, x, y0, 6.0)
    tb, tf = textbox(s4, x, y0 + 0.32, 6.05, 2.5)
    for i, parts in enumerate(rows):
        para(tf, i == 0, parts, size=10, bullet=True, space_after=5, line=1.10)

# ==================== SLIDE 5 - IMPACT AND BENEFITS =======================
s5 = S[4]
set_oval(s5)
set_title(s5, "IMPACT AND BENEFITS", 32)
drop(by_name(s5, "TextBox 8"))

BAND = 5.05
C1, C2, C3, CW = 0.40, 4.72, 9.04, 4.0
vline(s5, C2 - 0.20, BODY_TOP, BAND - 0.15)
vline(s5, C3 - 0.20, BODY_TOP, BAND - 0.15)
hline(s5, 0.40, 12.95, BAND - 0.05)

cols = [
    ("Direct Impact on Target Users", C1, [
        [("ISRO and NRSC analysts: ", True),
         ("plain-English querying replaces manual scene-by-scene inspection in a GIS "
          "desktop.", False)],
        [("Disaster response (NDRF/SDRF): ", True),
         ("rapid before-and-after assessment of flood, landslide and cyclone imagery - and "
          "because SAR sees through cloud, it works on the days optical imaging fails.",
          False)],
        [("Agriculture and land records: ", True),
         ("land-cover and change queries over Cartosat scenes without requiring "
          "remote-sensing expertise.", False)],
        [("Defence and border monitoring: ", True),
         ("combined optical and radar interpretation, with a written trace behind every "
          "conclusion.", False)],
    ]),
    ("Strategic Impact", C2, [
        [("Lowers the expertise barrier", True),
         (" to India's growing open Earth-observation archive, so more of it gets used.",
          False)],
        [("Makes AI output institutionally usable. ", True),
         ("An auditable trace is what allows a model's answer to enter an official decision "
          "record - the gap that blocks AI adoption in government today.", False)],
        [("Builds indigenous capability", True),
         (" in remote-sensing vision-language models rather than depending on foreign "
          "general-purpose systems, aligning with Atmanirbhar Bharat.", False)],
        [("All-weather capability. ", True),
         ("Radar interpretation keeps the system useful during monsoon and cloud cover.",
          False)],
    ]),
    ("Economic and Social Benefits", C3, [
        [("Removes the analyst bottleneck. ", True),
         ("Imagery volume grows continuously while trained interpreters do not; natural-"
          "language querying widens who can extract value from it.", False)],
        [("No procurement burden. ", True),
         ("Commodity hardware, open weights, open datasets - deployable without a GPU "
          "cluster or per-seat licensing.", False)],
        [("Faster disaster assessment", True),
         (" shortens the interval between an event and a resourcing decision.", False)],
        [("Data stays in-country. ", True),
         ("Self-hosted inference means sensitive imagery never leaves the institutional "
          "network.", False)],
    ]),
]
for title, x, rows in cols:
    section(s5, title, x, BODY_TOP, CW)
    tb, tf = textbox(s5, x, BODY_TOP + 0.32, CW, 3.3)
    for i, parts in enumerate(rows):
        para(tf, i == 0, parts, size=10, bullet=True, space_after=6, line=1.10)

stats = [
    ("47", "satellites already open on Bhoonidhi -\nthe archive this system unlocks"),
    ("4", "sensing capabilities in one system: VQA,\ncaptioning, change, optical+SAR fusion"),
    ("100%", "of model choices logged with a reason -\nzero opaque selections in the trace"),
    ("0", "GPU clusters required to deploy -\ncommodity hardware, self-hosted"),
]
sw = (12.95 - 0.40) / 4
for i, (num, desc) in enumerate(stats):
    x = 0.40 + i * sw
    tb, tf = textbox(s5, x, BAND + 0.18, sw - 0.15, 0.55)
    para(tf, True, [(num, True)], size=28, color=NAVY, align=PP_ALIGN.CENTER,
         space_after=0)
    tb, tf = textbox(s5, x, BAND + 0.78, sw - 0.15, 0.75)
    for j, ln_ in enumerate(desc.split("\n")):
        para(tf, j == 0, [(ln_, False)], size=9, align=PP_ALIGN.CENTER, space_after=0,
             line=1.10)

# ================= SLIDE 6 - RESEARCH AND REFERENCES ======================
s6 = S[5]
set_oval(s6)
set_title(s6, "RESEARCH AND REFERENCES", 32)
drop(by_name(s6, "TextBox 8"))

MID_X, ROW_Y = 6.68, 3.95
vline(s6, MID_X, BODY_TOP, BODY_BOT)
hline(s6, 0.40, MID_X - 0.05, ROW_Y - 0.10)
hline(s6, MID_X + 0.05, 12.95, ROW_Y - 0.10)

refquad = [
    ("Datasets Prescribed by the Problem Statement", 0.40, BODY_TOP, [
        [("BigEarthNet v2.0 / reBEN", True),
         (" - 549,488 co-registered Sentinel-1/2 patch pairs, 19-class CORINE labels. "
          "arXiv:2407.03653  \u00b7  bigearth.net  \u00b7  CDLA-Permissive-1.0", False)],
        [("VRSBench", True),
         (" - 29,614 images with captions, object references and question-answer pairs. "
          "arXiv:2406.12384  \u00b7  CC-BY-NC-4.0", False)],
        [("RSVQA (Lobry et al.)", True),
         (" - low- and high-resolution remote-sensing VQA. Zenodo 6344334 / 6344367  "
          "\u00b7  CC-BY-4.0", False)],
        [("CDVQA (Yuan et al., IEEE TGRS 2022)", True),
         (" - change-detection VQA over bi-temporal pairs; imagery from SECOND.", False)],
    ]),
    ("Existing Solutions & Competitive Analysis", MID_X + 0.25, BODY_TOP, [
        [("GeoChat (CVPR 2024)", True),
         (" - first grounded vision-language model for remote sensing; establishes that "
          "domain adaptation, not scale, drives accuracy on nadir imagery.", False)],
        [("TEOChat (ICLR 2025)", True),
         (" - temporal Earth-observation assistant. arXiv:2410.06234", False)],
        [("EarthDial (CVPR 2025)", True),
         (" - multi-sensory dialogue across RGB, SAR and multispectral inputs.", False)],
        [("GeoPixel (ICML 2025)", True),
         (" - pixel-level grounding at high resolution. arXiv:2501.13925", False)],
        [("Gap: ", True),
         ("all are single monolithic models. None expose an auditable execution trace or "
          "validate co-registration before answering.", False)],
    ]),
    ("Methods & Technology Benchmarking", 0.40, ROW_Y, [
        [("Domain-adaptive post-training for multimodal models", True),
         (" - arXiv:2411.19930. Evidence that a small domain-adapted model outperforms a "
          "larger generic one - the basis for our 3B-plus-specialist design.", False)],
        [("reBEN geographic split", True),
         (" - reduces spatial correlation between train and test, so reported accuracy is "
          "honest rather than inflated by neighbouring patches.", False)],
        [("SECOND (Yang et al.)", True),
         (" - 4,662 annotated bi-temporal pairs over Hangzhou, Chengdu and Shanghai. "
          "captain-whu.github.io/SCD/", False)],
        [("SAR radiometric practice", True),
         (" - decibel conversion with percentile clipping, standard before any model "
          "consumes backscatter.", False)],
    ]),
    ("Operational & Policy Context", MID_X + 0.25, ROW_Y, [
        [("Bhoonidhi (ISRO/NRSC)", True),
         (" - open dissemination of Earth-observation data from 47 satellites, including "
          "Cartosat-2S optical and EOS-04 (RISAT-1A) SAR products in GeoTIFF. "
          "bhoonidhi.nrsc.gov.in", False)],
        [("Indian Space Policy 2023", True),
         (" - mandates open dissemination of EO data, which makes an assistant over that "
          "archive worth building now.", False)],
        [("EOS-04 handbook (NRSC)", True),
         (" - product levels and radiometric conventions that determine correct SAR "
          "preprocessing.", False)],
        [("Evaluation protocol: ", True),
         ("accuracy on RSVQA and CDVQA test splits, plus a three-way ablation "
          "(optical-only, SAR-only, fused) on the BigEarthNet geographic test split.",
          False)],
    ]),
]
for title, x, y0, rows in refquad:
    section(s6, title, x, y0, 6.0)
    tb, tf = textbox(s6, x, y0 + 0.32, 6.05, 2.5)
    for i, parts in enumerate(rows):
        para(tf, i == 0, parts, size=9.5, bullet=True, space_after=5, line=1.10)

prs.save(OUT)
print("WROTE", OUT, os.path.getsize(OUT), "bytes")
